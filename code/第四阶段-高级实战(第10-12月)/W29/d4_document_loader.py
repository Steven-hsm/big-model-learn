"""
W29-D4: 多格式文档加载器
========================
功能:
  - 多格式文档加载 (PDF/Word/TXT/Markdown)
  - 文本提取
  - 元数据管理
"""

import os
import re
from dataclasses import dataclass, field
from typing import List, Dict, Optional
from datetime import datetime
from pathlib import Path


# ============================================================
# 1. 文档数据模型
# ============================================================

@dataclass
class DocumentMetadata:
    """文档元数据"""
    filename: str
    file_type: str           # pdf/docx/txt/md
    file_size: int           # 字节数
    created_at: str = ""
    page_count: int = 0
    title: str = ""
    author: str = ""
    extra: Dict[str, str] = field(default_factory=dict)


@dataclass
class LoadedDocument:
    """加载后的文档"""
    content: str
    metadata: DocumentMetadata
    source_path: str
    load_time_ms: float = 0.0
    char_count: int = 0
    word_count: int = 0

    def __post_init__(self):
        self.char_count = len(self.content)
        self.word_count = len(self.content.split())


# ============================================================
# 2. 各格式加载器
# ============================================================

class TextLoader:
    """纯文本加载器"""

    @staticmethod
    def load(file_path: str) -> LoadedDocument:
        """加载TXT文件"""
        import time
        start = time.time()

        path = Path(file_path)
        content = path.read_text(encoding="utf-8")
        metadata = DocumentMetadata(
            filename=path.name,
            file_type="txt",
            file_size=path.stat().st_size,
            created_at=datetime.fromtimestamp(
                path.stat().st_ctime
            ).isoformat(),
        )
        elapsed = (time.time() - start) * 1000

        return LoadedDocument(
            content=content,
            metadata=metadata,
            source_path=str(path),
            load_time_ms=elapsed,
        )


class MarkdownLoader:
    """Markdown加载器"""

    @staticmethod
    def load(file_path: str) -> LoadedDocument:
        """加载Markdown文件, 提取纯文本和结构信息"""
        import time
        start = time.time()

        path = Path(file_path)
        raw_content = path.read_text(encoding="utf-8")

        # 提取标题结构
        headers = re.findall(r'^(#{1,6})\s+(.+)$', raw_content, re.MULTILINE)

        # 移除Markdown语法, 提取纯文本
        text = raw_content
        text = re.sub(r'```[\s\S]*?```', '', text)  # 代码块
        text = re.sub(r'`[^`]+`', '', text)          # 行内代码
        text = re.sub(r'!\[.*?\]\(.*?\)', '', text)   # 图片
        text = re.sub(r'\[([^\]]+)\]\([^\)]+\)', r'\1', text)  # 链接
        text = re.sub(r'^#{1,6}\s+', '', text, flags=re.MULTILINE)  # 标题标记
        text = re.sub(r'[*_]{1,2}([^*_]+)[*_]{1,2}', r'\1', text)  # 粗体/斜体
        text = re.sub(r'^[-*+]\s+', '', text, flags=re.MULTILINE)  # 列表标记
        text = re.sub(r'^\d+\.\s+', '', text, flags=re.MULTILINE)  # 有序列表
        text = re.sub(r'^>\s+', '', text, flags=re.MULTILINE)       # 引用
        text = re.sub(r'\n{3,}', '\n\n', text)  # 多余空行

        metadata = DocumentMetadata(
            filename=path.name,
            file_type="md",
            file_size=path.stat().st_size,
            created_at=datetime.fromtimestamp(
                path.stat().st_ctime
            ).isoformat(),
            title=headers[0][1] if headers else path.stem,
            extra={"header_count": str(len(headers))},
        )
        elapsed = (time.time() - start) * 1000

        return LoadedDocument(
            content=text.strip(),
            metadata=metadata,
            source_path=str(path),
            load_time_ms=elapsed,
        )


class PDFLoader:
    """PDF加载器 (需要PyPDF2)"""

    @staticmethod
    def load(file_path: str) -> LoadedDocument:
        """加载PDF文件"""
        import time
        start = time.time()

        path = Path(file_path)
        pages_text = []
        metadata_extra = {}

        try:
            from PyPDF2 import PdfReader

            reader = PdfReader(str(path))
            metadata_extra["page_count"] = str(len(reader.pages))

            # 提取元数据
            if reader.metadata:
                if reader.metadata.title:
                    metadata_extra["title"] = reader.metadata.title
                if reader.metadata.author:
                    metadata_extra["author"] = reader.metadata.author

            # 逐页提取文本
            for i, page in enumerate(reader.pages):
                text = page.extract_text()
                if text:
                    pages_text.append(f"--- 第{i+1}页 ---\n{text}")

        except ImportError:
            # 没有PyPDF2, 使用模拟数据
            print("  [提示] 安装PyPDF2以支持PDF加载: pip install PyPDF2")
            pages_text = ["[需要安装PyPDF2来解析PDF内容]"]
            metadata_extra["note"] = "PyPDF2未安装"

        content = "\n\n".join(pages_text)
        elapsed = (time.time() - start) * 1000

        metadata = DocumentMetadata(
            filename=path.name,
            file_type="pdf",
            file_size=path.stat().st_size,
            created_at=datetime.fromtimestamp(
                path.stat().st_ctime
            ).isoformat(),
            page_count=int(metadata_extra.get("page_count", 0)),
            title=metadata_extra.get("title", ""),
            author=metadata_extra.get("author", ""),
            extra=metadata_extra,
        )

        return LoadedDocument(
            content=content,
            metadata=metadata,
            source_path=str(path),
            load_time_ms=elapsed,
        )


class DocxLoader:
    """Word文档加载器 (需要python-docx)"""

    @staticmethod
    def load(file_path: str) -> LoadedDocument:
        """加载Word文档"""
        import time
        start = time.time()

        path = Path(file_path)
        paragraphs = []
        metadata_extra = {}

        try:
            import docx

            doc = docx.Document(str(path))

            # 提取元数据
            core_props = doc.core_properties
            if core_props.title:
                metadata_extra["title"] = core_props.title
            if core_props.author:
                metadata_extra["author"] = core_props.author

            # 提取段落文本
            for para in doc.paragraphs:
                if para.text.strip():
                    paragraphs.append(para.text)

            # 提取表格文本
            for table in doc.tables:
                for row in table.rows:
                    row_text = " | ".join(
                        cell.text for cell in row.cells
                    )
                    paragraphs.append(row_text)

        except ImportError:
            print("  [提示] 安装python-docx以支持Word加载: pip install python-docx")
            paragraphs = ["[需要安装python-docx来解析Word文档]"]
            metadata_extra["note"] = "python-docx未安装"

        content = "\n\n".join(paragraphs)
        elapsed = (time.time() - start) * 1000

        metadata = DocumentMetadata(
            filename=path.name,
            file_type="docx",
            file_size=path.stat().st_size,
            created_at=datetime.fromtimestamp(
                path.stat().st_ctime
            ).isoformat(),
            title=metadata_extra.get("title", ""),
            author=metadata_extra.get("author", ""),
            extra=metadata_extra,
        )

        return LoadedDocument(
            content=content,
            metadata=metadata,
            source_path=str(path),
            load_time_ms=elapsed,
        )


# ============================================================
# 3. 统一文档加载器
# ============================================================

class DocumentLoader:
    """统一文档加载器"""

    # 支持的文件格式和对应的加载器
    LOADERS = {
        ".txt": TextLoader,
        ".md": MarkdownLoader,
        ".pdf": PDFLoader,
        ".docx": DocxLoader,
    }

    def __init__(self):
        self.supported_extensions = list(self.LOADERS.keys())

    def load(self, file_path: str) -> LoadedDocument:
        """根据文件类型选择加载器"""
        path = Path(file_path)

        if not path.exists():
            raise FileNotFoundError(f"文件不存在: {file_path}")

        ext = path.suffix.lower()
        if ext not in self.LOADERS:
            raise ValueError(
                f"不支持的文件格式: {ext}, "
                f"支持: {self.supported_extensions}"
            )

        loader = self.LOADERS[ext]
        return loader.load(file_path)

    def load_directory(self, dir_path: str) -> List[LoadedDocument]:
        """加载目录下所有支持的文件"""
        path = Path(dir_path)
        if not path.is_dir():
            raise NotADirectoryError(f"不是目录: {dir_path}")

        documents = []
        for file_path in path.rglob("*"):
            if file_path.suffix.lower() in self.LOADERS:
                try:
                    doc = self.load(str(file_path))
                    documents.append(doc)
                    print(f"  已加载: {file_path.name} "
                          f"({doc.char_count}字符, "
                          f"{doc.load_time_ms:.1f}ms)")
                except Exception as e:
                    print(f"  加载失败: {file_path.name} - {e}")

        return documents

    def get_file_info(self, file_path: str) -> Dict:
        """获取文件信息(不加载内容)"""
        path = Path(file_path)
        if not path.exists():
            return {"error": "文件不存在"}

        return {
            "name": path.name,
            "extension": path.suffix,
            "size_bytes": path.stat().st_size,
            "size_human": self._human_size(path.stat().st_size),
            "supported": path.suffix.lower() in self.LOADERS,
        }

    @staticmethod
    def _human_size(size: int) -> str:
        """人类可读的文件大小"""
        for unit in ["B", "KB", "MB", "GB"]:
            if size < 1024:
                return f"{size:.1f}{unit}"
            size /= 1024
        return f"{size:.1f}TB"


# ============================================================
# 4. 演示
# ============================================================

def demo():
    """演示文档加载器"""

    print("=" * 60)
    print("W29-D4: 多格式文档加载器")
    print("=" * 60)

    loader = DocumentLoader()
    print(f"\n支持的文件格式: {loader.supported_extensions}")

    # 创建测试文件
    import tempfile
    tmp_dir = tempfile.mkdtemp()

    # 测试TXT
    txt_file = os.path.join(tmp_dir, "sample.txt")
    with open(txt_file, "w", encoding="utf-8") as f:
        f.write("这是一个测试文档。\n\n包含多个段落。\n\n用于测试文本加载器。")

    # 测试Markdown
    md_file = os.path.join(tmp_dir, "sample.md")
    with open(md_file, "w", encoding="utf-8") as f:
        f.write("# 测试文档\n\n")
        f.write("## 第一节\n\n")
        f.write("这是**加粗文本**和*斜体文本*。\n\n")
        f.write("### 列表\n\n")
        f.write("- 项目一\n- 项目二\n- 项目三\n\n")
        f.write("[链接文本](https://example.com)\n")

    print("\n--- 加载TXT文件 ---")
    doc = loader.load(txt_file)
    print(f"  文件: {doc.metadata.filename}")
    print(f"  类型: {doc.metadata.file_type}")
    print(f"  大小: {doc.metadata.file_size} 字节")
    print(f"  字符数: {doc.char_count}")
    print(f"  词数: {doc.word_count}")
    print(f"  耗时: {doc.load_time_ms:.2f}ms")
    print(f"  内容预览: {doc.content[:80]}...")

    print("\n--- 加载Markdown文件 ---")
    doc = loader.load(md_file)
    print(f"  文件: {doc.metadata.filename}")
    print(f"  类型: {doc.metadata.file_type}")
    print(f"  标题: {doc.metadata.title}")
    print(f"  字符数: {doc.char_count}")
    print(f"  内容预览: {doc.content[:80]}...")

    print("\n--- 批量加载目录 ---")
    docs = loader.load_directory(tmp_dir)
    print(f"共加载 {len(docs)} 个文件")

    # 清理
    import shutil
    shutil.rmtree(tmp_dir)
    print("\n(临时文件已清理)")


# ============================================================
# 运行演示
# ============================================================

if __name__ == "__main__":
    demo()
